"""
MotorPH Analytical Dashboard  -  Milestone 2
============================================

Turns the Milestone 2 Draft (product category distribution chart) into a complete
analytical dashboard built from the MotorPH Products and Sales datasets.

What this script does
---------------------
1. LOAD      Reads the two preprocessed CSV files and stops with a clear message if a
             file or column is missing.
2. CLEAN     Validates and cleans both datasets with Pandas (types, blanks, duplicates,
             totals, unmatched products) and prints a data-quality report.
3. ANALYSE   Computes key metrics and summary tables (monthly, client type, payment
             method, product, category, inventory).
4. VISUALISE Draws the Matplotlib charts and assembles them into a four-section
             dashboard: MotorPH Overview, Sales Performance, Sales by Product Category,
             Inventory Section.
5. EXPORT    Saves the dashboard (PNG + PDF), the Milestone 2 chart, summary CSVs and a
             formula-driven Excel dashboard.

Run it
------
    python MotorPH_Analytical_Dashboard.py
    python MotorPH_Analytical_Dashboard.py --data-dir path/to/csvs --output-dir results
    python MotorPH_Analytical_Dashboard.py --no-show --no-excel

The two CSV files must sit in the data folder (default: the folder of this script):
    MotorPH_Products_Preprocessed.csv
    MotorPH_Sales_Preprocessed.csv

Requirements: pandas, matplotlib, openpyxl (Excel export only).
"""

# %%
# ----------------------------------------------------------------------------
# 0. IMPORTS AND CONFIGURATION
# ----------------------------------------------------------------------------
from __future__ import annotations

import argparse
import sys
import textwrap
from dataclasses import dataclass
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyBboxPatch
from matplotlib.ticker import FuncFormatter

# Colour palette: used by every chart so the dashboard looks consistent.
NAVY = "#12355B"
BLUE = "#2F6DB5"      # same blue as the Milestone 2 Draft chart
ORANGE = "#E8833A"    # highlight colour (largest / top item)
TEAL = "#2A9D8F"
GREY = "#6B7280"
LIGHT = "#F3F6FA"
BORDER = "#D5DEE9"
SERIES_COLOURS = [BLUE, ORANGE, TEAL, "#8E6BBF", "#C4C9D1"]

CURRENCY = "₱"        # MotorPH prices are in Philippine pesos (assumption stated in the report)

PRODUCT_COLS = ["Product ID Number", "Product Name", "Product Type", "Unit Price",
                "Date of Manufacturing", "Date of Acquisition"]
SALES_COLS = ["Date", "Client Type", "Product Name", "Unit Price", "Quantity",
              "Total", "Payment Method"]


@dataclass(frozen=True)
class Config:
    """All file locations and settings in one place."""
    data_dir: Path
    output_dir: Path
    products_file: str = "MotorPH_Products_Preprocessed.csv"
    sales_file: str = "MotorPH_Sales_Preprocessed.csv"
    show_plots: bool = True
    make_excel: bool = True

    @property
    def products_path(self) -> Path:
        return self.data_dir / self.products_file

    @property
    def sales_path(self) -> Path:
        return self.data_dir / self.sales_file


class DataError(Exception):
    """Raised when an input file cannot be used (missing, empty, wrong columns)."""


# %%
# ----------------------------------------------------------------------------
# 1. HELPERS
# ----------------------------------------------------------------------------
def money(value: float, symbol: str = CURRENCY) -> str:
    """Short currency text for labels: 4,459,496,800 -> ₱4.46B."""
    magnitude = abs(value)
    if magnitude >= 1e9:
        return f"{symbol}{value / 1e9:.2f}B"
    if magnitude >= 1e6:
        return f"{symbol}{value / 1e6:.1f}M"
    if magnitude >= 1e3:
        return f"{symbol}{value / 1e3:.0f}K"
    return f"{symbol}{value:,.0f}"


def money_axis(ax, axis: str = "y") -> None:
    """Format an axis with short currency labels."""
    formatter = FuncFormatter(lambda v, _: money(v))
    (ax.yaxis if axis == "y" else ax.xaxis).set_major_formatter(formatter)


def style_axes(ax, title: str, grid_axis: str = "y") -> None:
    """Apply the shared look: left-aligned bold title, no top/right frame, light grid."""
    ax.set_title(title, fontsize=13, fontweight="bold", loc="left", color=NAVY, pad=10)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis=grid_axis, alpha=0.25)
    ax.set_axisbelow(True)


# %%
# ----------------------------------------------------------------------------
# 2. LOAD AND CLEAN
# ----------------------------------------------------------------------------
def load_csv(path: Path, required: list[str]) -> pd.DataFrame:
    """Read a CSV file and confirm it is not empty and has the required columns."""
    if not path.exists():
        raise DataError(f"File not found: {path}\n"
                        f"Keep the CSV in the data folder or pass --data-dir.")
    data = pd.read_csv(path)
    if data.empty:
        raise DataError(f"{path.name} is empty.")
    missing = [c for c in required if c not in data.columns]
    if missing:
        raise DataError(f"{path.name} is missing columns {missing}. "
                        f"Found: {list(data.columns)}")
    return data


def clean_products(raw: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Clean the Products dataset (already preprocessed in Milestone 1; re-validated here)."""
    report = {"Product rows loaded": len(raw)}
    df = raw[PRODUCT_COLS].copy()

    for col in ["Product Name", "Product Type"]:
        df[col] = df[col].astype("string").str.strip().replace("", pd.NA)
    for col in ["Product ID Number", "Unit Price", "Date of Manufacturing", "Date of Acquisition"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    blank = df[["Product Name", "Product Type", "Unit Price"]].isna().any(axis=1)
    report["Products with blank name/category/price removed"] = int(blank.sum())
    df = df[~blank]

    report["Duplicate product rows removed"] = int(df.duplicated().sum())
    df = df.drop_duplicates()
    report["Duplicate Product IDs removed"] = int(df["Product ID Number"].duplicated().sum())
    df = df.drop_duplicates(subset="Product ID Number")

    df = df.astype({"Product ID Number": int, "Unit Price": int,
                    "Date of Manufacturing": int, "Date of Acquisition": int,
                    "Product Name": str, "Product Type": str})
    report["Product rows used"] = len(df)
    return df.reset_index(drop=True), report


def clean_sales(raw: pd.DataFrame, products: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Clean the Sales dataset and attach each sale's Product Type from the catalogue."""
    report = {"Sales rows loaded": len(raw)}
    df = raw[SALES_COLS].copy()

    # Text columns: trim spaces, treat blanks as missing, standardise client type casing.
    for col in ["Client Type", "Product Name", "Payment Method"]:
        df[col] = df[col].astype("string").str.strip().replace("", pd.NA)
    df["Client Type"] = df["Client Type"].str.title()

    # Correct data types: dates as datetime, numbers as numeric.
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    for col in ["Unit Price", "Quantity", "Total"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    # Missing values: rows without a usable date/product/price/quantity cannot be analysed.
    unusable = df[["Date", "Product Name", "Unit Price", "Quantity"]].isna().any(axis=1)
    unusable |= (df["Quantity"] <= 0) | (df["Unit Price"] <= 0)
    report["Rows removed (missing/invalid date, product, price or quantity)"] = int(unusable.sum())
    df = df[~unusable].copy()

    # Missing categories are kept (the sale is real) but labelled "Unknown".
    for col in ["Client Type", "Payment Method"]:
        report[f"Blank {col} labelled 'Unknown'"] = int(df[col].isna().sum())
        df[col] = df[col].fillna("Unknown")

    # Consistency: Total must equal Unit Price x Quantity.
    expected = df["Unit Price"] * df["Quantity"]
    mismatch = df["Total"].isna() | (df["Total"] != expected)
    report["Totals corrected (Unit Price x Quantity)"] = int(mismatch.sum())
    df["Total"] = expected

    report["Duplicate sales rows removed"] = int(df.duplicated().sum())
    df = df.drop_duplicates()

    # Integer types for counts/money, plain strings for text.
    df = df.astype({"Unit Price": int, "Quantity": int, "Total": int,
                    "Client Type": str, "Product Name": str, "Payment Method": str})

    # Integrate the two datasets: look up each sale's category in the product catalogue.
    lookup = products.set_index("Product Name")["Product Type"]
    df["Product Type"] = df["Product Name"].map(lookup).fillna("Uncategorized")
    report["Sales not matched to a catalogue product"] = int((df["Product Type"] == "Uncategorized").sum())

    # Sales prices should agree with the catalogue price (informational check only).
    catalogue_price = df["Product Name"].map(products.set_index("Product Name")["Unit Price"])
    report["Sales whose price differs from catalogue price"] = int((catalogue_price != df["Unit Price"]).sum())

    report["Sales rows used"] = len(df)
    return df.sort_values("Date").reset_index(drop=True), report


def print_report(title: str, report: dict) -> None:
    """Print a data-quality report in a readable block."""
    print(f"\n--- {title} ---")
    for key, value in report.items():
        print(f"{key}: {value:,}")


# %%
# ----------------------------------------------------------------------------
# 3. ANALYSIS: SUMMARY TABLES AND KEY METRICS
# ----------------------------------------------------------------------------
def group_summary(sales: pd.DataFrame, column: str) -> pd.DataFrame:
    """Revenue, units, orders and revenue share for each value of `column`."""
    table = (sales.groupby(column)
             .agg(Revenue=("Total", "sum"), Units=("Quantity", "sum"), Orders=("Total", "size"))
             .sort_values("Revenue", ascending=False).reset_index())
    table["Share (%)"] = (table["Revenue"] / table["Revenue"].sum() * 100).round(1)
    return table


def weekly_main_period(sales: pd.DataFrame, max_gap_days: int = 28) -> pd.DataFrame:
    """Weekly revenue for the main continuous sales period.

    Sales dates are split into periods wherever there is a gap longer than `max_gap_days`;
    the longest period is kept so isolated records do not stretch the time axis. Weeks that
    only partly fall inside that period are dropped so every plotted week is a full week.
    """
    days = pd.Series(sorted(sales["Date"].dt.normalize().unique()))
    period_id = (days.diff().dt.days > max_gap_days).cumsum()
    main = days.groupby(period_id).agg(["min", "max", "size"]).sort_values("size").iloc[-1]
    subset = sales[(sales["Date"] >= main["min"]) & (sales["Date"] <= main["max"])]
    weekly = subset.set_index("Date")["Total"].resample("W").sum()
    full = (weekly.index - pd.Timedelta(days=6) >= main["min"]) & (weekly.index <= main["max"])
    weekly = weekly[full]
    return weekly.rename("Revenue").rename_axis("Week").reset_index()


def build_tables(products: pd.DataFrame, sales: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Create every summary table used by the charts, the Excel file and the insights."""
    tables: dict[str, pd.DataFrame] = {}

    # Monthly: every calendar month between first and last sale (missing months become 0).
    months = pd.period_range(sales["Date"].min(), sales["Date"].max(), freq="M")
    monthly = (sales.groupby(sales["Date"].dt.to_period("M"))
               .agg(Revenue=("Total", "sum"), Units=("Quantity", "sum"), Orders=("Total", "size"))
               .reindex(months, fill_value=0))
    monthly.index = monthly.index.to_timestamp()
    tables["monthly"] = monthly.rename_axis("Month").reset_index()

    tables["weekly"] = weekly_main_period(sales)

    tables["client"] = group_summary(sales, "Client Type")
    tables["payment"] = group_summary(sales, "Payment Method")

    # Product performance: all catalogue products, including any with no sales.
    sold = sales.groupby("Product Name").agg(Units=("Quantity", "sum"), Revenue=("Total", "sum"))
    by_product = (products.merge(sold, left_on="Product Name", right_index=True, how="left")
                  .fillna({"Units": 0, "Revenue": 0}).astype({"Units": int, "Revenue": int}))
    tables["product"] = by_product.sort_values("Revenue", ascending=False).reset_index(drop=True)

    # Category performance: catalogue size next to sales results.
    counts = products["Product Type"].value_counts().rename("Products")
    cat_sales = sales.groupby("Product Type").agg(Units=("Quantity", "sum"), Revenue=("Total", "sum"))
    category = (pd.concat([counts, cat_sales], axis=1).fillna(0).astype(int)
                .rename_axis("Category").reset_index())
    category["Catalogue Share (%)"] = (category["Products"] / category["Products"].sum() * 100).round(1)
    category["Revenue Share (%)"] = (category["Revenue"] / category["Revenue"].sum() * 100).round(1)
    category["Revenue per Model"] = (category["Revenue"] / category["Products"]).round(0).astype(int)
    # Correctness check (same as the Milestone 2 Draft): counts must add up to the catalogue size.
    assert category["Products"].sum() == len(products), "Category counts do not match the catalogue."
    tables["category"] = category.sort_values("Products", ascending=False, kind="stable").reset_index(drop=True)

    # Inventory list: grouped by category, most expensive model first.
    inventory = tables["product"].sort_values(["Product Type", "Unit Price"],
                                              ascending=[True, False]).reset_index(drop=True)
    tables["inventory"] = inventory
    return tables


def compute_metrics(products: pd.DataFrame, sales: pd.DataFrame) -> dict:
    """Headline figures for the MotorPH Overview section."""
    revenue = int(sales["Total"].sum())
    orders = len(sales)
    return {
        "revenue": revenue,
        "units": int(sales["Quantity"].sum()),
        "orders": orders,
        "avg_order_value": revenue / orders,
        "n_products": len(products),
        "n_categories": products["Product Type"].nunique(),
        "avg_price": float(products["Unit Price"].mean()),
        "first_sale": sales["Date"].min(),
        "last_sale": sales["Date"].max(),
        "active_days": sales["Date"].nunique(),
    }


def build_insights(metrics: dict, tables: dict[str, pd.DataFrame]) -> list[str]:
    """Plain-language findings, generated from the data so they never go out of date."""
    category = tables["category"].sort_values("Revenue", ascending=False).reset_index(drop=True)
    product = tables["product"]
    client = tables["client"].set_index("Client Type")
    monthly = tables["monthly"]
    out: list[str] = []

    top = category.iloc[0]
    out.append(f"{top['Category']} is the top revenue category: {money(top['Revenue'])} "
               f"({top['Revenue Share (%)']:.1f}% of sales) from {top['Products']} of "
               f"{metrics['n_products']} models ({top['Catalogue Share (%)']:.0f}% of the catalogue).")

    top3 = category.head(3)
    out.append(f"The top 3 categories ({', '.join(top3['Category'])}) bring in "
               f"{top3['Revenue Share (%)'].sum():.0f}% of revenue, so sales are concentrated in a few segments.")

    efficient = category.sort_values("Revenue per Model", ascending=False).iloc[0]
    out.append(f"{efficient['Category']} earns the most per model ({money(efficient['Revenue per Model'])} "
               f"per product), which means the number of models in a category does not by itself predict sales.")

    best_rev = product.iloc[0]
    best_units = product.sort_values("Units", ascending=False).iloc[0]
    out.append(f"Best seller by revenue: {best_rev['Product Name']} ({money(best_rev['Revenue'])}). "
               f"Best seller by units: {best_units['Product Name']} ({best_units['Units']:,} units).")

    if {"Retail", "Wholesale"} <= set(client.index):
        out.append(f"Wholesale orders contribute {client.loc['Wholesale', 'Share (%)']:.1f}% of revenue and retail "
                   f"{client.loc['Retail', 'Share (%)']:.1f}%, so both channels matter to planning and stock.")

    empty = monthly[monthly["Orders"] == 0]["Month"].dt.strftime("%b %Y").tolist()
    sparse = monthly[(monthly["Orders"] > 0) &
                     (monthly["Orders"] < 0.05 * monthly.loc[monthly["Orders"] > 0, "Orders"].median())]["Month"].dt.strftime("%b %Y").tolist()
    notes = []
    if empty:
        notes.append(f"no sales records for {', '.join(empty)}")
    if sparse:
        notes.append(f"only a handful of records in {', '.join(sparse)}")
    if notes:
        out.append("Data coverage: sales run from "
                   f"{metrics['first_sale']:%d %b %Y} to {metrics['last_sale']:%d %b %Y}, with "
                   + " and ".join(notes) + ". Trend conclusions should be drawn from the months with full data.")
    return out


# %%
# ----------------------------------------------------------------------------
# 4. CHART FUNCTIONS (each draws on an Axes so it can be reused in the dashboard or alone)
# ----------------------------------------------------------------------------
def draw_category_count(ax, category: pd.DataFrame) -> None:
    """Milestone 2 Draft chart: number of products per category, largest on top."""
    ordered = category.sort_values("Products")
    bars = ax.barh(ordered["Category"], ordered["Products"], color=BLUE)
    offset = ordered["Products"].max() * 0.01
    for bar, count, pct in zip(bars, ordered["Products"], ordered["Catalogue Share (%)"]):
        ax.text(bar.get_width() + offset, bar.get_y() + bar.get_height() / 2,
                f"{count} ({pct:.0f}%)", va="center", fontsize=9)
    style_axes(ax, f"Products per Category (N = {int(ordered['Products'].sum())})", grid_axis="x")
    ax.set_xlabel("Number of products")
    ax.set_xlim(0, ordered["Products"].max() * 1.18)


def draw_category_revenue(ax, category: pd.DataFrame) -> None:
    """Revenue by category with each category's share of total revenue."""
    ordered = category.sort_values("Revenue")
    colours = [ORANGE if v == ordered["Revenue"].max() else BLUE for v in ordered["Revenue"]]
    bars = ax.barh(ordered["Category"], ordered["Revenue"], color=colours)
    offset = ordered["Revenue"].max() * 0.01
    for bar, rev, pct in zip(bars, ordered["Revenue"], ordered["Revenue Share (%)"]):
        ax.text(bar.get_width() + offset, bar.get_y() + bar.get_height() / 2,
                f"{money(rev)} ({pct:.1f}%)", va="center", fontsize=9)
    style_axes(ax, "Revenue by Product Category", grid_axis="x")
    ax.set_xlabel(f"Revenue ({CURRENCY})")
    money_axis(ax, "x")
    ax.set_xlim(0, ordered["Revenue"].max() * 1.30)


def draw_revenue_per_model(ax, category: pd.DataFrame) -> None:
    """Average revenue generated by each model in a category (sales efficiency)."""
    ordered = category.sort_values("Revenue per Model")
    colours = [ORANGE if v == ordered["Revenue per Model"].max() else TEAL for v in ordered["Revenue per Model"]]
    bars = ax.barh(ordered["Category"], ordered["Revenue per Model"], color=colours)
    offset = ordered["Revenue per Model"].max() * 0.01
    for bar, val in zip(bars, ordered["Revenue per Model"]):
        ax.text(bar.get_width() + offset, bar.get_y() + bar.get_height() / 2,
                money(val), va="center", fontsize=9)
    style_axes(ax, "Average Revenue per Model by Category", grid_axis="x")
    ax.set_xlabel(f"Revenue per model ({CURRENCY})")
    money_axis(ax, "x")
    ax.set_xlim(0, ordered["Revenue per Model"].max() * 1.2)


def draw_monthly_revenue(ax, monthly: pd.DataFrame) -> None:
    """Revenue per calendar month; months with no records are labelled, not hidden."""
    labels = monthly["Month"].dt.strftime("%b\n%Y")
    colours = [BLUE if n > 0 else BORDER for n in monthly["Orders"]]
    bars = ax.bar(labels, monthly["Revenue"], color=colours, width=0.65)
    top = max(monthly["Revenue"].max(), 1)
    for bar, rev, orders in zip(bars, monthly["Revenue"], monthly["Orders"]):
        if orders == 0:
            ax.text(bar.get_x() + bar.get_width() / 2, top * 0.02, "no\nrecords",
                    ha="center", va="bottom", fontsize=8, color=GREY, style="italic")
        else:
            ax.text(bar.get_x() + bar.get_width() / 2, rev + top * 0.015, money(rev),
                    ha="center", va="bottom", fontsize=9)
    style_axes(ax, "Monthly Revenue")
    ax.set_ylabel(f"Revenue ({CURRENCY})")
    money_axis(ax)
    ax.set_ylim(0, top * 1.15)


def draw_weekly_trend(ax, weekly: pd.DataFrame) -> None:
    """Weekly revenue across the main continuous sales period (full weeks only)."""
    ax.plot(weekly["Week"], weekly["Revenue"], color=BLUE, marker="o", markersize=5, linewidth=2)
    average = weekly["Revenue"].mean()
    ax.axhline(average, color=ORANGE, linestyle="--", linewidth=1.3)
    ax.text(weekly["Week"].iloc[0], average, f" weekly average {money(average)}",
            color=ORANGE, va="bottom", ha="left", fontsize=9)
    peak = weekly.loc[weekly["Revenue"].idxmax()]
    ax.annotate(f"Peak week: {money(peak['Revenue'])}", xy=(peak["Week"], peak["Revenue"]),
                xytext=(0, 14), textcoords="offset points", ha="center", fontsize=9, color=NAVY,
                arrowprops=dict(arrowstyle="-", color=GREY))
    start, end = weekly["Week"].iloc[0] - pd.Timedelta(days=6), weekly["Week"].iloc[-1]
    style_axes(ax, f"Weekly Revenue Trend ({start:%d %b} to {end:%d %b %Y}, full weeks)")
    ax.set_ylabel(f"Revenue per week ({CURRENCY})")
    ax.set_xlabel("Week ending")
    money_axis(ax)
    ax.set_ylim(0, weekly["Revenue"].max() * 1.22)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d %b"))
    ax.set_xticks(weekly["Week"])
    ax.tick_params(axis="x", rotation=45, labelsize=8.5)


def draw_group_bars(ax, table: pd.DataFrame, column: str, title: str) -> None:
    """Revenue by a grouping column (e.g. client type) with share labels."""
    colours = [SERIES_COLOURS[i % len(SERIES_COLOURS)] for i in range(len(table))]
    bars = ax.bar(table[column], table["Revenue"], color=colours, width=0.6)
    top = table["Revenue"].max()
    for bar, rev, pct in zip(bars, table["Revenue"], table["Share (%)"]):
        ax.text(bar.get_x() + bar.get_width() / 2, rev + top * 0.015,
                f"{money(rev)}\n{pct:.1f}%", ha="center", va="bottom", fontsize=9)
    style_axes(ax, title)
    ax.set_ylabel(f"Revenue ({CURRENCY})")
    money_axis(ax)
    ax.set_ylim(0, top * 1.25)


def draw_payment_donut(ax, payment: pd.DataFrame) -> None:
    """Donut chart of revenue by payment method (few categories, so a donut is readable)."""
    colours = [SERIES_COLOURS[i % len(SERIES_COLOURS)] for i in range(len(payment))]
    ax.pie(payment["Revenue"], colors=colours, startangle=90, counterclock=False,
           wedgeprops=dict(width=0.42, edgecolor="white"))
    ax.text(0, 0, f"{money(payment['Revenue'].sum())}\ntotal", ha="center", va="center",
            fontsize=11, fontweight="bold", color=NAVY)
    legend = [f"{m}  {money(r)} ({p:.1f}%)" for m, r, p in
              zip(payment["Payment Method"], payment["Revenue"], payment["Share (%)"])]
    ax.legend(ax.patches, legend, loc="upper center", bbox_to_anchor=(0.5, 0.02), frameon=False, fontsize=9.5)
    ax.set_title("Revenue by Payment Method", fontsize=13, fontweight="bold", loc="left", color=NAVY, pad=10)


def draw_top_products(ax, product: pd.DataFrame, n: int = 10) -> None:
    """Top n products by revenue, labelled with revenue and units sold."""
    top = product.head(n).iloc[::-1]
    bars = ax.barh(top["Product Name"], top["Revenue"], color=BLUE)
    offset = top["Revenue"].max() * 0.01
    for bar, rev, units in zip(bars, top["Revenue"], top["Units"]):
        ax.text(bar.get_width() + offset, bar.get_y() + bar.get_height() / 2,
                f"{money(rev)}  |  {units:,} units", va="center", fontsize=9)
    style_axes(ax, f"Top {n} Products by Revenue", grid_axis="x")
    ax.set_xlabel(f"Revenue ({CURRENCY})")
    money_axis(ax, "x")
    ax.set_xlim(0, top["Revenue"].max() * 1.32)


def draw_kpi_card(ax, label: str, value: str, sub: str) -> None:
    """One KPI tile for the Overview section."""
    ax.axis("off")
    ax.add_patch(FancyBboxPatch((0.02, 0.06), 0.96, 0.88, boxstyle="round,pad=0,rounding_size=0.06",
                                transform=ax.transAxes, facecolor=LIGHT, edgecolor=BORDER, linewidth=1.2))
    ax.text(0.5, 0.74, label, ha="center", va="center", fontsize=11, color=GREY, transform=ax.transAxes)
    ax.text(0.5, 0.46, value, ha="center", va="center", fontsize=22, fontweight="bold",
            color=NAVY, transform=ax.transAxes)
    ax.text(0.5, 0.18, sub, ha="center", va="center", fontsize=9.5, color=GREY, transform=ax.transAxes)


def draw_section_label(ax, text: str) -> None:
    """Section banner used between dashboard sections."""
    ax.axis("off")
    ax.add_patch(FancyBboxPatch((0, 0.1), 1, 0.8, boxstyle="round,pad=0,rounding_size=0.1",
                                transform=ax.transAxes, facecolor=NAVY, edgecolor="none"))
    ax.text(0.01, 0.5, text, ha="left", va="center", fontsize=16, fontweight="bold",
            color="white", transform=ax.transAxes)


def draw_insights(ax, insights: list[str]) -> None:
    """Key insights box: wrapped bullet points."""
    ax.axis("off")
    ax.add_patch(FancyBboxPatch((0, 0), 1, 1, boxstyle="round,pad=0,rounding_size=0.03",
                                transform=ax.transAxes, facecolor=LIGHT, edgecolor=BORDER))
    wrapped = ["\n".join(textwrap.wrap(line, width=175, initial_indent="•  ", subsequent_indent="    "))
               for line in insights]
    ax.text(0.015, 0.93, "\n\n".join(wrapped), ha="left", va="top", fontsize=11.5,
            color="#1F2937", transform=ax.transAxes, linespacing=1.35)


# %%
# ----------------------------------------------------------------------------
# 5. DASHBOARD ASSEMBLY
# ----------------------------------------------------------------------------
def build_dashboard_page(metrics: dict, tables: dict[str, pd.DataFrame], insights: list[str]):
    """Page 1: MotorPH Overview, Sales Performance and Sales by Product Category."""
    fig = plt.figure(figsize=(20, 28))
    heights = [1.0, 1.3, 0.45, 4.0, 4.6, 0.45, 6.2, 0.45, 3.2]
    outer = fig.add_gridspec(len(heights), 1, height_ratios=heights, hspace=0.42,
                             left=0.065, right=0.975, top=0.985, bottom=0.015)

    def row(index: int, ratios: list[float], wspace: float):
        """Sub-grid for one dashboard row: one axes per chart with its own spacing."""
        sub = outer[index].subgridspec(1, len(ratios), width_ratios=ratios, wspace=wspace)
        return [fig.add_subplot(sub[0, i]) for i in range(len(ratios))]

    # Header
    header = row(0, [1], 0)[0]
    header.axis("off")
    header.text(-0.055, 0.62, "MotorPH Analytical Dashboard", fontsize=30, fontweight="bold",
                color=NAVY, transform=header.transAxes)
    header.text(-0.055, 0.12, f"Product lineup and sales performance  |  Sales records from "
                f"{metrics['first_sale']:%d %b %Y} to {metrics['last_sale']:%d %b %Y}  |  Amounts in Philippine pesos ({CURRENCY})",
                fontsize=12.5, color=GREY, transform=header.transAxes)

    # Section A: MotorPH Overview (KPI tiles)
    cards = [
        ("Total Revenue", money(metrics["revenue"]), f"{metrics['orders']:,} sales orders"),
        ("Total Units Sold", f"{metrics['units']:,}", f"{metrics['active_days']} days with sales"),
        ("Average Order Value", money(metrics["avg_order_value"]), "revenue per order"),
        ("Products in Catalogue", f"{metrics['n_products']}", f"{metrics['n_categories']} categories"),
        ("Average Unit Price", money(metrics["avg_price"]), "across catalogue"),
        ("Avg Units per Order", f"{metrics['units'] / metrics['orders']:.1f}", "units per sales order"),
    ]
    for ax, (label, value, sub) in zip(row(1, [1] * 6, 0.04), cards):
        draw_kpi_card(ax, label, value, sub)

    # Section B: Sales Performance
    draw_section_label(row(2, [1], 0)[0], "Sales Performance")
    ax_month, ax_week = row(3, [5, 7], 0.22)
    draw_monthly_revenue(ax_month, tables["monthly"])
    draw_weekly_trend(ax_week, tables["weekly"])
    ax_client, ax_pay, ax_top = row(4, [3, 3.4, 6.6], 0.5)
    draw_group_bars(ax_client, tables["client"], "Client Type", "Revenue by Client Type")
    draw_payment_donut(ax_pay, tables["payment"])
    draw_top_products(ax_top, tables["product"])

    # Section C: Sales by Product Category (includes the Milestone 2 Draft chart)
    draw_section_label(row(5, [1], 0)[0], "Sales by Product Category")
    ax_count, ax_rev, ax_per = row(6, [1, 1, 1], 0.55)
    draw_category_count(ax_count, tables["category"])
    draw_category_revenue(ax_rev, tables["category"])
    draw_revenue_per_model(ax_per, tables["category"])

    # Key insights
    draw_section_label(row(7, [1], 0)[0], "Key Insights")
    draw_insights(row(8, [1], 0)[0], insights)
    return fig


def build_inventory_page(inventory: pd.DataFrame):
    """Page 2: Inventory Section - a structured list of products, categories and pricing."""
    fig = plt.figure(figsize=(20, 19))
    gs = fig.add_gridspec(2, 1, height_ratios=[0.5, 12], hspace=0.02, left=0.03, right=0.97, top=0.985, bottom=0.01)
    draw_section_label(fig.add_subplot(gs[0]), "Inventory Section: Products, Categories and Pricing")

    ax = fig.add_subplot(gs[1])
    ax.axis("off")
    columns = ["ID", "Product Name", "Category", f"Unit Price ({CURRENCY})",
               "Mfg Year", "Acquired", "Units Sold", f"Revenue ({CURRENCY})"]
    cells = [[int(r["Product ID Number"]), r["Product Name"], r["Product Type"], f"{r['Unit Price']:,}",
              int(r["Date of Manufacturing"]), int(r["Date of Acquisition"]),
              f"{r['Units']:,}", f"{r['Revenue']:,}"] for _, r in inventory.iterrows()]
    table = ax.table(cellText=cells, colLabels=columns, loc="upper center", cellLoc="left",
                     colWidths=[0.05, 0.25, 0.14, 0.13, 0.09, 0.09, 0.10, 0.15])
    table.auto_set_font_size(False)
    table.set_fontsize(10)
    table.scale(1, 1.62)

    group_ids = inventory["Product Type"].astype("category").cat.codes.tolist()
    numeric_cols = {3, 4, 5, 6, 7}
    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor(BORDER)
        if row == 0:
            cell.set_facecolor(NAVY)
            cell.set_text_props(color="white", fontweight="bold")
        else:
            cell.set_facecolor(LIGHT if group_ids[row - 1] % 2 == 0 else "white")
        if col in numeric_cols and row > 0:
            cell.set_text_props(ha="right")
    return fig


def save_figure(fig, path: Path, dpi: int = 130) -> None:
    fig.savefig(path, dpi=dpi, facecolor="white")
    print(f"Saved {path}")


def show_saved_image(path: Path) -> None:
    """Open a saved dashboard image in a window.

    The dashboard is drawn at a large fixed size (20 x 28 inches). Showing the live figure in a
    smaller window squeezes the layout and makes labels collide, so the saved image is displayed
    instead: it looks exactly like the PNG and the toolbar zoom/pan tools can be used to read details.
    """
    image = plt.imread(path)
    height, width = image.shape[:2]
    fig = plt.figure(figsize=(7, 7 * height / width), num=path.stem)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.imshow(image)
    ax.axis("off")


def export_dashboard(cfg: Config, metrics, tables, insights) -> None:
    """Save dashboard pages (PNG + one PDF), the Milestone 2 chart and the summary CSVs."""
    out = cfg.output_dir
    page1 = build_dashboard_page(metrics, tables, insights)
    page2 = build_inventory_page(tables["inventory"])
    save_figure(page1, out / "MotorPH_Dashboard.png")
    save_figure(page2, out / "MotorPH_Dashboard_Inventory.png")
    with PdfPages(out / "MotorPH_Dashboard.pdf") as pdf:
        pdf.savefig(page1, facecolor="white")
        pdf.savefig(page2, facecolor="white")
    print(f"Saved {out / 'MotorPH_Dashboard.pdf'}")

    # The Milestone 2 Draft chart as a stand-alone file
    draft, ax = plt.subplots(figsize=(10, max(4, 0.55 * len(tables["category"]) + 1.5)), layout="constrained")
    draw_category_count(ax, tables["category"])
    save_figure(draft, out / "product_distribution_by_category.png", dpi=200)

    tables_dir = out / "tables"
    tables_dir.mkdir(exist_ok=True)
    for name, table in tables.items():
        table.to_csv(tables_dir / f"{name}_summary.csv", index=False)
    print(f"Saved summary tables in {tables_dir}")

    plt.close("all")
    if cfg.show_plots:
        show_saved_image(out / "MotorPH_Dashboard.png")
        show_saved_image(out / "MotorPH_Dashboard_Inventory.png")
        plt.show()


# %%
# ----------------------------------------------------------------------------
# 6. EXCEL DASHBOARD (live formulas: change the data sheets and everything recalculates)
# ----------------------------------------------------------------------------
def export_excel(path: Path, products: pd.DataFrame, sales: pd.DataFrame,
                 tables: dict[str, pd.DataFrame], insights: list[str], metrics: dict) -> None:
    """Build an Excel workbook: Dashboard, Summary Tables, Products and Sales Data sheets."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    font = "Arial"
    navy, blue, light = NAVY.lstrip("#"), BLUE.lstrip("#"), LIGHT.lstrip("#")
    thin = Side(style="thin", color=BORDER.lstrip("#"))
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    peso = f'"{CURRENCY}"#,##0'
    peso_m = f'"{CURRENCY}"#,##0.0,,"M"'

    def style(cell, bold=False, size=10, color="000000", fill=None, align=None, fmt=None, border=False, wrap=False):
        cell.font = Font(name=font, bold=bold, size=size, color=color)
        if fill:
            cell.fill = PatternFill("solid", start_color=fill)
        if align or wrap:
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if fmt:
            cell.number_format = fmt
        if border:
            cell.border = box

    def header_row(ws, row, col, labels):
        for i, text in enumerate(labels):
            c = ws.cell(row=row, column=col + i, value=text)
            style(c, bold=True, color="FFFFFF", fill=navy, align="center", border=True)

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    summ = wb.create_sheet("Summary Tables")
    prod_ws = wb.create_sheet("Products")
    sales_ws = wb.create_sheet("Sales Data")

    # ---- Sales Data sheet (cleaned data; Product Type and Total are formulas) ----
    n_sales, n_prod = len(sales), len(products)
    header_row(sales_ws, 1, 1, SALES_COLS + ["Product Type"])
    for i, r in enumerate(sales.itertuples(index=False), start=2):
        sales_ws.cell(i, 1, r.Date.to_pydatetime()).number_format = "yyyy-mm-dd"
        sales_ws.cell(i, 2, r[1])
        sales_ws.cell(i, 3, r[2])
        sales_ws.cell(i, 4, r[3]).number_format = peso
        sales_ws.cell(i, 5, r[4])
        sales_ws.cell(i, 6, f"=D{i}*E{i}").number_format = peso
        sales_ws.cell(i, 7, r[6])
        sales_ws.cell(i, 8, f"=INDEX(Products!$C$2:$C${n_prod + 1},MATCH(C{i},Products!$B$2:$B${n_prod + 1},0))")
        for c in range(1, 9):
            sales_ws.cell(i, c).font = Font(name=font, size=10)
    for col, w in zip("ABCDEFGH", [12, 13, 28, 14, 10, 16, 15, 15]):
        sales_ws.column_dimensions[col].width = w
    sales_ws.freeze_panes = "A2"
    sales_ws.auto_filter.ref = f"A1:H{n_sales + 1}"

    def sd(col):  # absolute range on the Sales Data sheet
        return f"'Sales Data'!${col}$2:${col}${n_sales + 1}"

    # ---- Products sheet (catalogue + live units/revenue) ----
    inventory = tables["inventory"]
    header_row(prod_ws, 1, 1, ["Product ID Number", "Product Name", "Product Type", "Unit Price",
                               "Date of Manufacturing", "Date of Acquisition", "Units Sold", "Revenue"])
    for i, r in enumerate(inventory.itertuples(index=False), start=2):
        values = [r[0], r[1], r[2], r[3], r[4], r[5],
                  f"=SUMIFS({sd('E')},{sd('C')},B{i})", f"=SUMIFS({sd('F')},{sd('C')},B{i})"]
        for c, v in enumerate(values, start=1):
            cell = prod_ws.cell(i, c, v)
            style(cell, fmt=peso if c in (4, 8) else ("#,##0" if c == 7 else None))
    for col, w in zip("ABCDEFGH", [14, 30, 16, 14, 20, 20, 12, 16]):
        prod_ws.column_dimensions[col].width = w
    prod_ws.freeze_panes = "A2"
    prod_ws.auto_filter.ref = f"A1:H{n_prod + 1}"
    prng = lambda col: f"Products!${col}$2:${col}${n_prod + 1}"

    # ---- Summary Tables sheet (every chart reads from here) ----
    summ["A1"] = "Summary Tables (all values are formulas on the Sales Data and Products sheets)"
    style(summ["A1"], bold=True, size=14, color=navy)

    monthly = tables["monthly"]
    summ["A3"] = "Monthly revenue"
    style(summ["A3"], bold=True, size=11, color=navy)
    header_row(summ, 4, 1, ["Month", "Revenue", "Units"])
    m0, m1 = 5, 4 + len(monthly)
    for i, month in enumerate(monthly["Month"], start=m0):
        summ.cell(i, 1, month.to_pydatetime()).number_format = "mmm yyyy"
        summ.cell(i, 2, f'=SUMIFS({sd("F")},{sd("A")},">="&A{i},{sd("A")},"<"&EDATE(A{i},1))').number_format = peso_m
        summ.cell(i, 3, f'=SUMIFS({sd("E")},{sd("A")},">="&A{i},{sd("A")},"<"&EDATE(A{i},1))').number_format = "#,##0"

    def group_block(col, title, key_col, names):
        """Revenue / share table for client type or payment method."""
        letter = lambda k: chr(ord("A") + col - 1 + k)
        summ.cell(3, col, title)
        style(summ.cell(3, col), bold=True, size=11, color=navy)
        header_row(summ, 4, col, [title.split(" by ")[-1].title(), "Revenue", "Share"])
        first, last = 5, 4 + len(names)
        for i, name in enumerate(names, start=first):
            summ.cell(i, col, name)
            summ.cell(i, col + 1, f"=SUMIFS({sd('F')},{sd(key_col)},{letter(0)}{i})").number_format = peso_m
            summ.cell(i, col + 2, f"={letter(1)}{i}/SUM({letter(1)}${first}:{letter(1)}${last})").number_format = "0.0%"
        return first, last

    c0, c1 = group_block(5, "Revenue by client type", "B", tables["client"]["Client Type"].tolist())
    p0, p1 = group_block(9, "Revenue by payment method", "G", tables["payment"]["Payment Method"].tolist())

    def top_block(col, title, sort_col, n=10):
        """Top-n products; the ranking is fixed at build time, the values stay live."""
        top = tables["product"].sort_values(sort_col, ascending=False).head(n)
        summ.cell(3, col, title)
        style(summ.cell(3, col), bold=True, size=11, color=navy)
        header_row(summ, 4, col, ["Product", "Revenue", "Units"])
        for i, name in enumerate(top["Product Name"], start=5):
            summ.cell(i, col, name)
            summ.cell(i, col + 1, f"=SUMIFS({sd('F')},{sd('C')},{chr(ord('A') + col - 1)}{i})").number_format = peso_m
            summ.cell(i, col + 2, f"=SUMIFS({sd('E')},{sd('C')},{chr(ord('A') + col - 1)}{i})").number_format = "#,##0"
        return 5, 4 + n

    t0, t1 = top_block(13, "Top 10 products by revenue", "Revenue")
    u0, u1 = top_block(17, "Top 10 products by units sold", "Units")

    # Category tables
    category_count = tables["category"]
    category_rev = category_count.sort_values("Revenue", ascending=False)
    summ["A20"] = "Products per category (Milestone 2 Draft chart)"
    style(summ["A20"], bold=True, size=11, color=navy)
    header_row(summ, 21, 1, ["Category", "Products", "% of catalogue"])
    k0, k1 = 22, 21 + len(category_count)
    for i, name in enumerate(category_count["Category"], start=k0):
        summ.cell(i, 1, name)
        summ.cell(i, 2, f"=COUNTIF({prng('C')},A{i})").number_format = "0"
        summ.cell(i, 3, f"=B{i}/SUM(B${k0}:B${k1})").number_format = "0.0%"

    summ["E20"] = "Category sales performance"
    style(summ["E20"], bold=True, size=11, color=navy)
    header_row(summ, 21, 5, ["Category", "Products", "Units Sold", "Revenue", "Revenue Share", "Revenue per Model"])
    r0, r1 = 22, 21 + len(category_rev)
    for i, name in enumerate(category_rev["Category"], start=r0):
        summ.cell(i, 5, name)
        summ.cell(i, 6, f"=COUNTIF({prng('C')},E{i})").number_format = "0"
        summ.cell(i, 7, f"=SUMIFS({sd('E')},{sd('H')},E{i})").number_format = "#,##0"
        summ.cell(i, 8, f"=SUMIFS({sd('F')},{sd('H')},E{i})").number_format = peso_m
        summ.cell(i, 9, f"=H{i}/SUM(H${r0}:H${r1})").number_format = "0.0%"
        summ.cell(i, 10, f"=IFERROR(H{i}/F{i},0)").number_format = peso_m
    for col in range(1, 21):
        summ.column_dimensions[chr(ord("A") + col - 1)].width = 17
    for col in ("M", "Q"):
        summ.column_dimensions[col].width = 28
    for row in summ.iter_rows(min_row=3, max_row=r1):
        for cell in row:
            if cell.value is not None and cell.font.name != font:
                style(cell)

    # ---- Dashboard sheet ----
    dash.sheet_view.showGridLines = False
    for col in range(1, 19):
        dash.column_dimensions[chr(ord("A") + col - 1)].width = 11.5

    def banner(row, text):
        dash.merge_cells(start_row=row, start_column=1, end_row=row, end_column=18)
        dash.cell(row, 1, text)
        style(dash.cell(row, 1), bold=True, size=14, color="FFFFFF", fill=navy, align="left")
        dash.row_dimensions[row].height = 26

    dash.merge_cells("A1:R1")
    dash["A1"] = "MotorPH Analytical Dashboard"
    style(dash["A1"], bold=True, size=24, color=navy, align="left")
    dash.row_dimensions[1].height = 38
    dash.merge_cells("A2:R2")
    dash["A2"] = (f"Product lineup and sales performance | Sales records {metrics['first_sale']:%d %b %Y} to "
                  f"{metrics['last_sale']:%d %b %Y} | Amounts in Philippine pesos ({CURRENCY})")
    style(dash["A2"], size=10, color=GREY.lstrip("#"), align="left")

    # Overview KPI cards: label / value (formula) / note
    banner(4, "MotorPH Overview")
    cards = [
        ("Total Revenue", f"=SUM({sd('F')})", peso, "all sales orders"),
        ("Total Units Sold", f"=SUM({sd('E')})", "#,##0", "all sales orders"),
        ("Sales Orders", f"=COUNTA({sd('C')})", "#,##0", "rows in Sales Data"),
        ("Average Order Value", f"=SUM({sd('F')})/COUNTA({sd('C')})", peso, "revenue per order"),
        ("Products in Catalogue", f"=COUNTA({prng('B')})", "0", "see Inventory section"),
        ("Average Unit Price", f"=AVERAGE({prng('D')})", peso, "across catalogue"),
    ]
    for i, (label, formula, fmt, note) in enumerate(cards):
        c = 1 + i * 3
        for r, content in zip((5, 6, 7), (label, formula, note)):
            dash.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 2)
            dash.cell(r, c, content)
            for k in range(3):
                dash.cell(r, c + k).fill = PatternFill("solid", start_color=light)
        style(dash.cell(5, c), size=10, color=GREY.lstrip("#"), fill=light, align="center")
        style(dash.cell(6, c), bold=True, size=18, color=navy, fill=light, align="center", fmt=fmt)
        style(dash.cell(7, c), size=9, color=GREY.lstrip("#"), fill=light, align="center")
        dash.row_dimensions[6].height = 32

    def bar_chart(title, cats, vals, anchor, horizontal=False, width=12.6, height=7.6, fmt=None, colour=blue, labels=True):
        chart = BarChart()
        chart.type = "bar" if horizontal else "col"
        chart.title = title
        chart.style = 10
        chart.legend = None
        chart.add_data(vals, titles_from_data=True)
        chart.set_categories(cats)
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        if fmt:
            chart.y_axis.number_format = fmt
        if horizontal:
            chart.x_axis.scaling.orientation = "maxMin"
            chart.y_axis.crosses = "max"   # keep the value axis at the bottom
        chart.series[0].graphicalProperties.solidFill = colour
        chart.series[0].graphicalProperties.line.solidFill = colour
        if labels:
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
            for flag in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
                setattr(chart.dataLabels, flag, False)
            if fmt:
                chart.dataLabels.numFmt = fmt
        chart.width, chart.height = width, height
        dash.add_chart(chart, anchor)

    # Sales Performance
    banner(9, "Sales Performance")
    bar_chart("Monthly Revenue", Reference(summ, min_col=1, min_row=m0, max_row=m1),
              Reference(summ, min_col=2, min_row=4, max_row=m1), "A10", fmt=peso_m)
    bar_chart("Revenue by Client Type", Reference(summ, min_col=5, min_row=c0, max_row=c1),
              Reference(summ, min_col=6, min_row=4, max_row=c1), "G10", fmt=peso_m, colour=BLUE.lstrip("#"))
    pie = PieChart()
    pie.title = "Revenue by Payment Method"
    pie.add_data(Reference(summ, min_col=10, min_row=4, max_row=p1), titles_from_data=True)
    pie.set_categories(Reference(summ, min_col=9, min_row=p0, max_row=p1))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    for flag in ("showVal", "showSerName", "showCatName", "showLegendKey"):
        setattr(pie.dataLabels, flag, False)
    pie.width, pie.height = 12.6, 7.6
    dash.add_chart(pie, "M10")
    bar_chart("Top 10 Products by Revenue", Reference(summ, min_col=13, min_row=t0, max_row=t1),
              Reference(summ, min_col=14, min_row=4, max_row=t1), "A26", horizontal=True,
              width=19.0, height=8.6, fmt=peso_m)
    bar_chart("Top 10 Products by Units Sold", Reference(summ, min_col=17, min_row=u0, max_row=u1),
              Reference(summ, min_col=19, min_row=4, max_row=u1), "J26", horizontal=True,
              width=19.0, height=8.6, fmt="#,##0", colour=TEAL.lstrip("#"))

    # Sales by Product Category
    banner(44, "Sales by Product Category")
    bar_chart("Products per Category (Milestone 2 Draft)", Reference(summ, min_col=1, min_row=k0, max_row=k1),
              Reference(summ, min_col=2, min_row=21, max_row=k1), "A45", horizontal=True,
              width=12.6, height=10.5, fmt="0")
    bar_chart("Revenue by Product Category", Reference(summ, min_col=5, min_row=r0, max_row=r1),
              Reference(summ, min_col=8, min_row=21, max_row=r1), "G45", horizontal=True,
              width=12.6, height=10.5, fmt=peso_m)
    bar_chart("Average Revenue per Model", Reference(summ, min_col=5, min_row=r0, max_row=r1),
              Reference(summ, min_col=10, min_row=21, max_row=r1), "M45", horizontal=True,
              width=12.6, height=10.5, fmt=peso_m, colour=TEAL.lstrip("#"))

    # Key insights (text generated by the Python script from the data at build time)
    banner(67, "Key Insights (generated from the data when the workbook was built)")
    for i, line in enumerate(insights, start=68):
        dash.merge_cells(start_row=i, start_column=1, end_row=i, end_column=18)
        dash.cell(i, 1, f"•  {line}")
        style(dash.cell(i, 1), size=10, wrap=True, align="left")
        dash.row_dimensions[i].height = 30

    # Inventory section: structured list read live from the Products sheet
    inv_banner = 68 + len(insights) + 1
    banner(inv_banner, "Inventory Section: Products, Categories and Pricing")
    head = inv_banner + 1
    spans = [("ID", 1, 1), ("Product Name", 2, 4), ("Category", 5, 6), ("Unit Price", 7, 8),
             ("Mfg Year", 9, 9), ("Acquired", 10, 10), ("Units Sold", 11, 11), ("Revenue", 12, 13)]
    for label, a, b in spans:
        if b > a:
            dash.merge_cells(start_row=head, start_column=a, end_row=head, end_column=b)
        style(dash.cell(head, a), bold=True, color="FFFFFF", fill=navy, align="center")
        dash.cell(head, a, label)
        for k in range(a, b + 1):
            dash.cell(head, k).fill = PatternFill("solid", start_color=navy)
    src_cols = ["A", "B", "C", "D", "E", "F", "G", "H"]
    fmts = [None, None, None, peso, "0", "0", "#,##0", peso]
    for n in range(n_prod):
        row = head + 1 + n
        for (label, a, b), src, fmt in zip(spans, src_cols, fmts):
            if b > a:
                dash.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
            dash.cell(row, a, f"=Products!{src}{n + 2}")
            band = light if n % 2 == 0 else "FFFFFF"
            for k in range(a, b + 1):
                dash.cell(row, k).fill = PatternFill("solid", start_color=band)
                dash.cell(row, k).border = box
            style(dash.cell(row, a), size=10, fill=band, fmt=fmt, border=True,
                  align="left" if label in ("Product Name", "Category") else "right")

    dash.sheet_properties.tabColor = navy
    wb.save(path)
    print(f"Saved {path}")


# %%
# ----------------------------------------------------------------------------
# 7. MAIN PROGRAM
# ----------------------------------------------------------------------------
def default_data_dir() -> Path:
    """Folder of this script; falls back to the working folder in notebooks / interactive cells."""
    try:
        return Path(__file__).resolve().parent
    except NameError:
        return Path.cwd()


def parse_args(argv=None) -> Config:
    parser = argparse.ArgumentParser(description="Build the MotorPH analytical dashboard.")
    parser.add_argument("--data-dir", type=Path, default=default_data_dir(),
                        help="folder containing the two MotorPH CSV files (default: script folder)")
    parser.add_argument("--output-dir", type=Path, default=None,
                        help="where to save results (default: <data-dir>/dashboard_output)")
    parser.add_argument("--no-show", action="store_true", help="save files without opening chart windows")
    parser.add_argument("--no-excel", action="store_true", help="skip the Excel workbook")
    args, _ = parser.parse_known_args(argv)  # parse_known_args keeps Jupyter/VS Code cells working
    out = args.output_dir or args.data_dir / "dashboard_output"
    return Config(data_dir=args.data_dir, output_dir=out,
                  show_plots=not args.no_show, make_excel=not args.no_excel)


def main(argv=None) -> None:
    cfg = parse_args(argv)
    if not cfg.show_plots:
        plt.switch_backend("Agg")
    cfg.output_dir.mkdir(parents=True, exist_ok=True)

    # 1-2. Load and clean
    try:
        products, product_report = clean_products(load_csv(cfg.products_path, PRODUCT_COLS))
        sales, sales_report = clean_sales(load_csv(cfg.sales_path, SALES_COLS), products)
    except DataError as err:
        sys.exit(f"ERROR: {err}")
    print_report("Products data-quality report", product_report)
    print_report("Sales data-quality report", sales_report)

    # 3. Analyse
    tables = build_tables(products, sales)
    metrics = compute_metrics(products, sales)
    insights = build_insights(metrics, tables)
    print("\n--- Key metrics ---")
    print(f"Total revenue: {CURRENCY}{metrics['revenue']:,}   Units sold: {metrics['units']:,}   "
          f"Orders: {metrics['orders']:,}   Avg order value: {CURRENCY}{metrics['avg_order_value']:,.0f}")
    print(f"Products: {metrics['n_products']}   Categories: {metrics['n_categories']}")
    print("\n--- Key insights ---")
    for line in insights:
        print(f"- {line}")

    # 4-5. Visualise and export
    print()
    export_dashboard(cfg, metrics, tables, insights)
    if cfg.make_excel:
        export_excel(cfg.output_dir / "MotorPH_Dashboard.xlsx", products, sales, tables, insights, metrics)


if __name__ == "__main__":
    main()